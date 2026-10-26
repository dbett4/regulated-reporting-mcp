"""Durable, local-only workbook review workbench.

Workbook contents are treated as hostile input and are only read.  Reviews are
dispositions on evidence, never an accounting approval or financial resolution.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import re
import sqlite3
import zipfile
from contextlib import contextmanager
from copy import deepcopy
from datetime import date, datetime, timezone
from decimal import Decimal, localcontext
from pathlib import Path
from typing import Any

from defusedxml import ElementTree
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from workiva_mcp.reporting_math import SCALES, number, signed_sum

LIMITS = {
    "compressed_bytes": 10_000_000,
    "expanded_bytes": 50_000_000,
    "zip_entries": 500,
    "sheets": 50,
    "cells": 200_000,
    "workbooks": 25,
    "rules": 100,
    "supporting_per_rule": 50,
    "xml_elements": 500_000,
    "xml_depth": 32,
    "xml_text_bytes": 20_000_000,
    "xml_attribute_bytes": 5_000_000,
    "relationships": 2_000,
}
_CELL = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")
_TOLERANCE = re.compile(r"^\d{1,18}(?:\.\d{1,6})?$")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _canonical(value: Any) -> bytes:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=True).encode()


def _display(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _inspect_xlsx(content: bytes) -> None:
    if not content or len(content) > LIMITS["compressed_bytes"]:
        raise ValueError("Workbook exceeds compressed size limit")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            filenames = [item.filename for item in infos]
            if len(filenames) != len(set(filenames)):
                raise ValueError("Duplicate archive entries are not accepted")
            if len(infos) > LIMITS["zip_entries"]:
                raise ValueError("Workbook exceeds archive entry limit")
            if sum(item.file_size for item in infos) > LIMITS["expanded_bytes"]:
                raise ValueError("Workbook exceeds expanded size limit")
            names = set(filenames)
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ValueError("Malformed XLSX archive")
            lowered = {name.lower() for name in names}
            if any("externallinks/" in name or name.endswith("vbaproject.bin") for name in lowered):
                raise ValueError("External links and macros are not accepted")
            # Bound styles, strings, relationships and metadata too: ZIP size
            # limits alone do not bound openpyxl's Python-object amplification.
            elements = text_bytes = attribute_bytes = relationships = 0
            for name in sorted(names):
                if not name.lower().endswith((".xml", ".rels")):
                    continue
                depth = 0
                relationship_ids = set()
                with archive.open(name) as part:
                    for event, element in ElementTree.iterparse(part, events=("start", "end")):
                        if event == "start":
                            depth += 1
                            elements += 1
                            attribute_bytes += sum(
                                len(k.encode()) + len(v.encode()) for k, v in element.attrib.items()
                            )
                            if (
                                elements > LIMITS["xml_elements"]
                                or depth > LIMITS["xml_depth"]
                                or len(element.attrib) > 32
                                or attribute_bytes > LIMITS["xml_attribute_bytes"]
                            ):
                                raise ValueError("Workbook exceeds XML complexity limit")
                            if element.tag.rsplit("}", 1)[-1] == "Relationship":
                                relationships += 1
                                relationship_id = element.attrib.get("Id")
                                if (
                                    not relationship_id
                                    or relationship_id in relationship_ids
                                    or relationships > LIMITS["relationships"]
                                ):
                                    raise ValueError("Invalid or excessive workbook relationships")
                                relationship_ids.add(relationship_id)
                        else:
                            depth -= 1
                            text_bytes += len((element.text or "").encode()) + len(
                                (element.tail or "").encode()
                            )
                            if text_bytes > LIMITS["xml_text_bytes"]:
                                raise ValueError("Workbook exceeds XML text limit")
                            element.clear()
            sheet_xml = [
                name
                for name in names
                if name.startswith("xl/worksheets/") and name.endswith(".xml")
            ]
            if len(sheet_xml) > LIMITS["sheets"]:
                raise ValueError("Workbook exceeds sheet limit")
            cell_count = 0
            for name in sheet_xml:
                maximum_row = maximum_column = 0
                for _, element in ElementTree.iterparse(archive.open(name), events=("end",)):
                    local_name = element.tag.rsplit("}", 1)[-1]
                    if local_name == "dimension":
                        end = element.attrib.get("ref", "").split(":")[-1]
                        if end:
                            row, column = coordinate_to_tuple(end.replace("$", ""))
                            maximum_row, maximum_column = (
                                max(maximum_row, row),
                                max(maximum_column, column),
                            )
                    elif local_name == "row":
                        row = int(element.attrib.get("r", "0"))
                        if not 1 <= row <= 1_048_576:
                            raise ValueError("Invalid worksheet row")
                        maximum_row = max(maximum_row, row)
                    elif local_name == "c":
                        cell_count += 1
                        ref = element.attrib.get("r")
                        if not ref or not _CELL.fullmatch(ref.replace("$", "")):
                            raise ValueError("Invalid worksheet cell coordinate")
                        row, column = coordinate_to_tuple(ref.replace("$", ""))
                        maximum_row, maximum_column = (
                            max(maximum_row, row),
                            max(maximum_column, column),
                        )
                    element.clear()
                    if (
                        cell_count > LIMITS["cells"]
                        or maximum_row * maximum_column > LIMITS["cells"]
                    ):
                        raise ValueError("Workbook exceeds worksheet iteration limit")
            if any(item.flag_bits & 1 for item in infos):
                raise ValueError("Encrypted archive entries are not accepted")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Malformed XLSX archive") from exc


def _snapshot(content: bytes) -> dict[str, dict[str, dict[str, Any]]]:
    _inspect_xlsx(content)
    raw_book = value_book = None
    try:
        raw_book = load_workbook(
            io.BytesIO(content), data_only=False, read_only=True, keep_links=False
        )
        value_book = load_workbook(
            io.BytesIO(content), data_only=True, read_only=True, keep_links=False
        )
    except Exception as exc:
        if raw_book is not None:
            raw_book.close()
        raise ValueError("Malformed or unsupported XLSX workbook") from exc
    try:
        if (
            len(raw_book.sheetnames) > LIMITS["sheets"]
            or raw_book.sheetnames != value_book.sheetnames
        ):
            raise ValueError("Invalid workbook sheets")
        result: dict[str, dict[str, dict[str, Any]]] = {}
        count = 0
        for name in raw_book.sheetnames:
            raw_sheet, value_sheet = raw_book[name], value_book[name]
            if not raw_sheet._worksheet_path.startswith("xl/worksheets/"):
                raise ValueError("Unsupported worksheet location")
            if (raw_sheet.max_row or 0) * (raw_sheet.max_column or 0) > LIMITS["cells"]:
                raise ValueError("Workbook exceeds worksheet iteration limit")
            cells: dict[str, dict[str, Any]] = {}
            for raw_row, value_row in zip(
                raw_sheet.iter_rows(), value_sheet.iter_rows(), strict=True
            ):
                for raw_cell, value_cell in zip(raw_row, value_row, strict=True):
                    if raw_cell.value is None and value_cell.value is None:
                        continue
                    count += 1
                    if count > LIMITS["cells"]:
                        raise ValueError("Workbook exceeds cell limit")
                    raw = _display(raw_cell.value)
                    calculated = _display(value_cell.value)
                    item = {"raw_value": raw, "value": calculated}
                    if isinstance(raw, str) and raw.startswith("=") and calculated is None:
                        item["error"] = "formula cached result unavailable"
                    elif raw_cell.data_type == "e" or value_cell.data_type == "e":
                        item["error"] = "spreadsheet error"
                    cells[raw_cell.coordinate] = item
            result[name] = cells
        return result
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Malformed or unsupported XLSX workbook") from exc
    finally:
        if raw_book is not None:
            raw_book.close()
        if value_book is not None:
            value_book.close()


def demo_workbook(version: int) -> bytes:
    """Return a synthetic City of Riverton XLSX fixture."""
    if version not in (1, 2):
        raise ValueError("Demo version must be 1 or 2")
    book = Workbook()
    report = book.active
    report.title = "Report"
    support = book.create_sheet("Support")
    report["A1"] = "City of Riverton"
    report["B2"] = 1_000_000 if version == 1 else 1_005_000
    support["B2"], support["B3"] = 600_000, 400_000
    report["B4"], support["B4"] = 250_000, 250_000  # unaffected control
    report["B6"] = "=SUM(Support!B2:B3)" if version == 1 else "=Support!B2+Support!B3"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _demo_rules() -> list[dict]:
    return [
        {
            "id": "cash",
            "description": "Cash agrees to supporting balances",
            "reported": {"sheet": "Report", "cell": "B2"},
            "supporting": [{"sheet": "Support", "cell": "B2"}, {"sheet": "Support", "cell": "B3"}],
        },
        {
            "id": "unaffected",
            "description": "Unaffected clean total",
            "reported": {"sheet": "Report", "cell": "B4"},
            "supporting": [{"sheet": "Support", "cell": "B4"}],
        },
        {
            "id": "formula",
            "description": "Raw formula evidence",
            "reported": {"sheet": "Report", "cell": "B6"},
            "supporting": [{"sheet": "Support", "cell": "B2"}, {"sheet": "Support", "cell": "B3"}],
        },
        {
            "id": "missing",
            "description": "Missing evidence remains visible",
            "reported": {"sheet": "Report", "cell": "B8"},
            "supporting": [{"sheet": "Support", "cell": "B8"}],
        },
    ]


class Workbench:
    def __init__(self, root: Path):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True, mode=0o700)
        os.chmod(self.root, 0o700)
        self.path = self.root / "workbench.sqlite3"
        fd = os.open(self.path, os.O_CREAT | os.O_WRONLY, 0o600)
        os.close(fd)
        os.chmod(self.path, 0o600)
        with self._db() as db:
            db.execute(
                "CREATE TABLE IF NOT EXISTS workspace (singleton INTEGER PRIMARY KEY CHECK(singleton=1), state TEXT NOT NULL)"
            )
            db.execute(
                "CREATE TABLE IF NOT EXISTS blobs (id TEXT PRIMARY KEY, content BLOB NOT NULL, snapshot TEXT NOT NULL)"
            )
            if db.execute("SELECT 1 FROM workspace").fetchone() is None:
                db.execute("INSERT INTO workspace VALUES (1, ?)", (json.dumps(self._empty()),))
            db.commit()

    @contextmanager
    def _db(self):
        db = sqlite3.connect(self.path, timeout=10)
        db.row_factory = sqlite3.Row
        try:
            yield db
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    @staticmethod
    def _empty() -> dict:
        return {
            "revision": 0,
            "workbooks": [],
            "active_version": None,
            "rules": [],
            "mappings_confirmed": False,
            "mapping_confirmation": None,
            "checks_current": False,
            "demo_version": None,
            "demo_lineage": None,
            "findings": [],
            "changes": [],
            "events": [],
            "summary": {"passed": 0, "failed": 0, "not_evaluated": 0, "reopened": 0},
        }

    def state(self) -> dict:
        with self._db() as db:
            state = json.loads(
                db.execute("SELECT state FROM workspace WHERE singleton=1").fetchone()[0]
            )
        for key, value in (
            ("mapping_confirmation", None),
            ("checks_current", False),
            ("demo_version", None),
            ("demo_lineage", None),
        ):
            state.setdefault(key, value)
        return state

    def preview(
        self,
        expected_revision: int,
        sheet: str | None = None,
        start_row: int = 1,
        start_column: int = 1,
    ) -> dict:
        """Return one read-only page from the retained active workbook snapshot."""
        for value, name, maximum in (
            (expected_revision, "expected revision", None),
            (start_row, "start row", 1_048_576),
            (start_column, "start column", 16_384),
        ):
            if (
                isinstance(value, bool)
                or not isinstance(value, int)
                or (maximum is not None and not 1 <= value <= maximum)
            ):
                raise ValueError(f"Invalid {name}")
        if sheet is not None and not isinstance(sheet, str):
            raise ValueError("Invalid sheet")

        with self._db() as db:
            db.execute("BEGIN")
            state = json.loads(
                db.execute("SELECT state FROM workspace WHERE singleton=1").fetchone()[0]
            )
            if state["revision"] != expected_revision:
                raise ValueError("Stale workspace revision")
            if state["active_version"] is None:
                raise ValueError("No active workbook")
            row = db.execute(
                "SELECT snapshot FROM blobs WHERE id=?", (state["active_version"],)
            ).fetchone()
            if row is None:
                raise ValueError("No active workbook")
            snapshot = json.loads(row[0])

            sheets = list(snapshot)
            selected = sheets[0] if sheet is None else sheet
            if selected not in snapshot:
                raise ValueError("Unknown sheet")

            # Preserve the fixed 20x8 contract at the bottom and right Excel
            # edges by moving the page origin inward rather than overflowing.
            page_row = min(start_row, 1_048_576 - 20 + 1)
            page_column = min(start_column, 16_384 - 8 + 1)
            cells = snapshot[selected]
            total_rows = total_columns = 1
            for coordinate in cells:
                cell_row, cell_column = coordinate_to_tuple(coordinate)
                total_rows = max(total_rows, cell_row)
                total_columns = max(total_columns, cell_column)

            rows = []
            for cell_row in range(page_row, page_row + 20):
                output_row = []
                for cell_column in range(page_column, page_column + 8):
                    coordinate = f"{get_column_letter(cell_column)}{cell_row}"
                    observed = cells.get(coordinate)
                    output_row.append(
                        None
                        if observed is None
                        else {
                            "cell": coordinate,
                            "raw_value": observed.get("raw_value"),
                            "value": observed.get("value"),
                            "error": observed.get("error"),
                        }
                    )
                rows.append(output_row)

            result = {
                "revision": state["revision"],
                "active_version": state["active_version"],
                "sheets": sheets,
                "sheet": selected,
                "start_row": page_row,
                "start_column": page_column,
                "total_rows": total_rows,
                "total_columns": total_columns,
                "rows": rows,
            }
            db.commit()
            return result

    def _mutate(self, expected: int, operation) -> dict:
        if isinstance(expected, bool) or not isinstance(expected, int):
            raise ValueError("Invalid expected revision")
        with self._db() as db:
            db.execute("BEGIN IMMEDIATE")
            state = json.loads(
                db.execute("SELECT state FROM workspace WHERE singleton=1").fetchone()[0]
            )
            for key, value in (
                ("mapping_confirmation", None),
                ("checks_current", False),
                ("demo_version", None),
                ("demo_lineage", None),
            ):
                state.setdefault(key, value)
            if state["revision"] != expected:
                raise ValueError("Stale workspace revision")
            operation(state, db)
            state["revision"] += 1
            db.execute("UPDATE workspace SET state=? WHERE singleton=1", (json.dumps(state),))
            db.commit()
        return state

    @staticmethod
    def _event(state: dict, kind: str, detail: dict) -> None:
        state["events"].append({"at": _now(), "kind": kind, "detail": detail})

    def import_workbook(
        self, content: bytes, filename: str, label: str, expected_revision: int
    ) -> dict:
        if not isinstance(content, bytes):
            raise ValueError("Workbook content must be bytes")
        if (
            not isinstance(filename, str)
            or not filename.strip()
            or len(filename) > 255
            or not isinstance(label, str)
            or not label.strip()
            or len(label) > 100
        ):
            raise ValueError("Invalid filename or label")
        snapshot = _snapshot(content)
        digest = hashlib.sha256(content).hexdigest()

        def operation(state, db):
            if len(state["workbooks"]) >= LIMITS["workbooks"]:
                raise ValueError("Workbook limit reached")
            if any(item["sha256"] == digest for item in state["workbooks"]):
                raise ValueError("Workbook already imported")
            workbook_id = digest
            before = self._active_snapshot(db, state)
            rerun = state["mappings_confirmed"]
            db.execute(
                "INSERT INTO blobs VALUES (?, ?, ?)", (workbook_id, content, json.dumps(snapshot))
            )
            state["workbooks"].append(
                {
                    "id": workbook_id,
                    "label": label,
                    "filename": filename,
                    "sha256": digest,
                    "imported_at": _now(),
                }
            )
            state["active_version"] = workbook_id
            state["changes"] = self._changes(before, snapshot)
            state["mappings_confirmed"] = False
            state["mapping_confirmation"] = None
            state["checks_current"] = False
            state["demo_version"] = None
            state["demo_lineage"] = None
            self._event(
                state,
                "workbook_imported",
                {
                    "id": workbook_id,
                    "label": label,
                    "sha256": digest,
                    "changes": len(state["changes"]),
                },
            )
            if rerun:
                self._run(state, snapshot)
                state["checks_current"] = False

        return self._mutate(expected_revision, operation)

    @staticmethod
    def _active_snapshot(db, state) -> dict:
        if state["active_version"] is None:
            return {}
        row = db.execute(
            "SELECT snapshot FROM blobs WHERE id=?", (state["active_version"],)
        ).fetchone()
        return json.loads(row[0])

    @staticmethod
    def _changes(before: dict, after: dict) -> list[dict]:
        changes = []
        for sheet in sorted(set(before) | set(after)):
            for cell in sorted(set(before.get(sheet, {})) | set(after.get(sheet, {}))):
                old, new = before.get(sheet, {}).get(cell), after.get(sheet, {}).get(cell)
                if old != new:
                    changes.append({"sheet": sheet, "cell": cell, "before": old, "after": new})
        return changes

    def save_rules(self, rules: list[dict], expected_revision: int) -> dict:
        normalized = self._validate_rules(rules)

        def operation(state, db):
            old_rules = {r["id"]: r for r in state["rules"]}
            new_rules = {r["id"]: r for r in normalized}
            state["rules"] = normalized
            state["mappings_confirmed"] = False
            state["mapping_confirmation"] = None
            state["checks_current"] = False
            for finding in state["findings"]:
                if finding.get("review_status") == "reviewed" and old_rules.get(
                    finding["rule_id"]
                ) != new_rules.get(finding["rule_id"]):
                    finding["review_status"] = "reopened"
            state["summary"]["reopened"] = sum(
                f["review_status"] == "reopened" for f in state["findings"]
            )
            self._event(
                state,
                "rules_saved",
                {"rules": deepcopy(normalized), "rules_digest": self._rules_digest(normalized)},
            )

        return self._mutate(expected_revision, operation)

    @staticmethod
    def _validate_rules(rules: list[dict]) -> list[dict]:
        if not isinstance(rules, list) or not 1 <= len(rules) <= LIMITS["rules"]:
            raise ValueError("Provide 1 to 100 rules")
        output = []
        ids = set()
        for supplied in rules:
            if not isinstance(supplied, dict) or set(supplied) - {
                "id",
                "description",
                "reported",
                "supporting",
                "unit",
                "decimal_places",
                "tolerance",
            }:
                raise ValueError("Invalid rule fields")
            rule = deepcopy(supplied)
            if (
                not isinstance(rule.get("id"), str)
                or not 1 <= len(rule["id"]) <= 100
                or rule["id"] in ids
            ):
                raise ValueError("Rule IDs must be unique and at most 100 characters")
            ids.add(rule["id"])
            if (
                not isinstance(rule.get("description"), str)
                or not rule["description"].strip()
                or len(rule["description"]) > 500
            ):
                raise ValueError("Invalid rule description")
            if (
                not isinstance(rule.get("supporting"), list)
                or not 1 <= len(rule["supporting"]) <= LIMITS["supporting_per_rule"]
            ):
                raise ValueError("Invalid supporting mappings")
            refs = [rule.get("reported"), *rule["supporting"]]
            for index, ref in enumerate(refs):
                allowed = {"sheet", "cell", "scale"} | ({"sign"} if index else set())
                if (
                    not isinstance(ref, dict)
                    or set(ref) - allowed
                    or not isinstance(ref.get("sheet"), str)
                    or not ref["sheet"].strip()
                    or len(ref["sheet"]) > 100
                    or not isinstance(ref.get("cell"), str)
                    or not _CELL.fullmatch(ref["cell"])
                ):
                    raise ValueError("Invalid cell mapping")
                row, column = coordinate_to_tuple(ref["cell"])
                if row > 1_048_576 or column > 16_384:
                    raise ValueError("Cell mapping exceeds XLSX bounds")
                ref.setdefault("scale", "units")
                if not isinstance(ref["scale"], str) or ref["scale"] not in SCALES:
                    raise ValueError("Invalid scale")
                if index:
                    ref.setdefault("sign", "add")
                    if ref["sign"] not in ("add", "subtract"):
                        raise ValueError("Invalid sign")
            rule.setdefault("unit", "USD")
            rule.setdefault("decimal_places", 2)
            rule.setdefault("tolerance", "0.00")
            if (
                not isinstance(rule["unit"], str)
                or not rule["unit"].strip()
                or len(rule["unit"]) > 30
                or isinstance(rule["decimal_places"], bool)
                or not isinstance(rule["decimal_places"], int)
                or not 0 <= rule["decimal_places"] <= 6
            ):
                raise ValueError("Invalid rule units or precision")
            if not isinstance(rule["tolerance"], str) or not _TOLERANCE.fullmatch(
                rule["tolerance"]
            ):
                raise ValueError("Invalid tolerance")
            output.append(rule)
        return output

    def confirm_mappings(self, reviewer: str, expected_revision: int) -> dict:
        if not isinstance(reviewer, str) or not reviewer.strip() or len(reviewer) > 100:
            raise ValueError("Invalid reviewer")

        def operation(state, db):
            if not state["rules"] or state["active_version"] is None:
                raise ValueError("Rules and workbook are required")
            snapshot = self._active_snapshot(db, state)
            # Confirmation cannot bless references to nonexistent worksheets.
            if any(
                ref["sheet"] not in snapshot
                for rule in state["rules"]
                for ref in [rule["reported"], *rule["supporting"]]
            ):
                raise ValueError("Mapping references an unknown sheet")
            state["mappings_confirmed"] = True
            confirmation = {
                "reviewer": reviewer,
                "at": _now(),
                "workbook_version": state["active_version"],
                "rules_digest": self._rules_digest(state["rules"]),
            }
            state["mapping_confirmation"] = confirmation
            state["checks_current"] = False
            self._event(state, "mappings_confirmed", confirmation)

        return self._mutate(expected_revision, operation)

    def run_checks(self, expected_revision: int) -> dict:
        def operation(state, db):
            if not state["mappings_confirmed"] or state["active_version"] is None:
                raise ValueError("Mappings must be confirmed")
            self._run(state, self._active_snapshot(db, state))

        return self._mutate(expected_revision, operation)

    def _run(self, state: dict, snapshot: dict) -> None:
        old = {finding["rule_id"]: finding for finding in state["findings"]}
        findings = [self._finding(rule, snapshot) for rule in state["rules"]]
        for finding in findings:
            finding["workbook_version"] = state["active_version"]
            finding["rules_digest"] = self._rules_digest(state["rules"])
            prior = old.get(finding["rule_id"])
            if not prior or prior.get("review_status") not in ("reviewed", "reopened"):
                continue
            if prior["evidence_digest"] == finding["evidence_digest"]:
                finding["review_status"] = prior["review_status"]
                finding["review"] = prior["review"]
            else:
                finding["review_status"] = "reopened"
                finding["review"] = prior["review"]
        state["findings"] = findings
        state["summary"] = {
            status: sum(f["status"] == status for f in findings)
            for status in ("passed", "failed", "not_evaluated")
        }
        state["summary"]["reopened"] = sum(f["review_status"] == "reopened" for f in findings)
        state["checks_current"] = True
        self._event(
            state, "checks_run", {"summary": state["summary"], "version": state["active_version"]}
        )

    @staticmethod
    def _finding(rule: dict, snapshot: dict) -> dict:
        evidence = []
        refs = [
            ("reported", rule["reported"]),
            *(("supporting", ref) for ref in rule["supporting"]),
        ]
        with localcontext() as context:
            context.prec = 80
            for role, ref in refs:
                observed = snapshot.get(ref["sheet"], {}).get(ref["cell"])
                item = {"role": role, "source": deepcopy(ref), "raw_value": None, "value": None}
                if observed is None:
                    item["error"] = "source cell is missing"
                else:
                    item.update(observed)
                    if "error" not in item:
                        try:
                            item["normalized_value"] = str(
                                number(item["value"]) * SCALES[ref["scale"]]
                            )
                        except ValueError as exc:
                            item["error"] = str(exc)
                evidence.append(item)
            finding = {
                "rule_id": rule["id"],
                "description": rule["description"],
                "status": "not_evaluated",
                "reported": None,
                "expected": None,
                "difference": None,
                "unit": rule["unit"],
                "reason": None,
                "evidence": evidence,
                "review_status": "unreviewed",
                "review": None,
            }
            if any("error" in item for item in evidence):
                finding["reason"] = "One or more source cells could not be evaluated"
            else:
                reported, expected, difference = signed_sum(
                    evidence[0]["normalized_value"],
                    (
                        (item["normalized_value"], ref["sign"])
                        for ref, item in zip(rule["supporting"], evidence[1:], strict=True)
                    ),
                    rule["decimal_places"],
                )
                finding.update(
                    status="passed" if abs(difference) <= Decimal(rule["tolerance"]) else "failed",
                    reported=str(reported),
                    expected=str(expected),
                    difference=str(difference),
                )
            finding["evidence_digest"] = hashlib.sha256(
                _canonical({"rule": rule, "evidence": evidence})
            ).hexdigest()
            return finding

    @staticmethod
    def _rules_digest(rules: list[dict]) -> str:
        return hashlib.sha256(_canonical(rules)).hexdigest()

    def review_finding(
        self,
        rule_id: str,
        disposition: str,
        note: str,
        reviewer: str,
        evidence_digest: str,
        expected_revision: int,
    ) -> dict:
        if (
            disposition not in ("noted", "needs_follow_up")
            or not isinstance(note, str)
            or not note.strip()
            or len(note) > 2000
            or not isinstance(reviewer, str)
            or not reviewer.strip()
            or len(reviewer) > 100
        ):
            raise ValueError("Invalid review disposition, note, or reviewer")

        def operation(state, db):
            confirmation = state.get("mapping_confirmation") or {}
            if (
                not state["mappings_confirmed"]
                or not state["checks_current"]
                or confirmation.get("workbook_version") != state["active_version"]
                or confirmation.get("rules_digest") != self._rules_digest(state["rules"])
            ):
                raise ValueError("Mappings and checks are not current")
            finding = next((item for item in state["findings"] if item["rule_id"] == rule_id), None)
            if finding is None:
                raise ValueError("Unknown finding")
            if finding["evidence_digest"] != evidence_digest:
                raise ValueError("Stale finding evidence")
            review = {"reviewer": reviewer, "note": note, "disposition": disposition, "at": _now()}
            finding.update(review_status="reviewed", review=review)
            self._event(
                state,
                "finding_reviewed",
                {
                    "rule_id": rule_id,
                    "workbook_version": state["active_version"],
                    "evidence_digest": evidence_digest,
                    "review": review,
                    "meaning": "reviewer disposition only; not approved or resolved",
                },
            )
            state["summary"]["reopened"] = sum(
                f["review_status"] == "reopened" for f in state["findings"]
            )

        return self._mutate(expected_revision, operation)

    def load_demo(self, version: int, expected_revision: int) -> dict:
        if version not in (1, 2):
            raise ValueError("Demo version must be 1 or 2")
        content = demo_workbook(version)
        snapshot = _snapshot(content)
        digest = hashlib.sha256(content).hexdigest()

        def operation(state, db):
            if version == 1 and (state["workbooks"] or state["rules"]):
                raise ValueError("Demo v1 requires an empty workspace")
            if version == 2 and (
                state.get("demo_lineage") != "riverton" or state.get("demo_version") != 1
            ):
                raise ValueError("Demo v2 requires demo v1")
            before = self._active_snapshot(db, state)
            rerun = state["mappings_confirmed"]
            db.execute(
                "INSERT INTO blobs VALUES (?, ?, ?)", (digest, content, json.dumps(snapshot))
            )
            state["workbooks"].append(
                {
                    "id": digest,
                    "label": f"City of Riverton v{version}",
                    "filename": f"riverton-v{version}.xlsx",
                    "sha256": digest,
                    "imported_at": _now(),
                }
            )
            state["active_version"] = digest
            state["changes"] = self._changes(before, snapshot)
            state["mappings_confirmed"] = False
            state["mapping_confirmation"] = None
            state["checks_current"] = False
            state["demo_lineage"] = "riverton"
            state["demo_version"] = version
            if version == 1:
                state["rules"] = self._validate_rules(_demo_rules())
                self._event(
                    state,
                    "rules_saved",
                    {
                        "rules": deepcopy(state["rules"]),
                        "rules_digest": self._rules_digest(state["rules"]),
                    },
                )
            self._event(
                state,
                "workbook_imported",
                {
                    "id": digest,
                    "label": f"City of Riverton v{version}",
                    "sha256": digest,
                    "changes": len(state["changes"]),
                    "demo_version": version,
                },
            )
            if rerun:
                self._run(state, snapshot)
                state["checks_current"] = False

        return self._mutate(expected_revision, operation)

    def export_packet(self) -> dict:
        state = self.state()
        with self._db() as db:
            versions = []
            for metadata in state["workbooks"]:
                row = db.execute(
                    "SELECT snapshot FROM blobs WHERE id=?", (metadata["id"],)
                ).fetchone()
                versions.append({**metadata, "evidence": json.loads(row[0])})
        packet = {
            "schema": "workbench.packet.v1",
            "exported_at": _now(),
            "status": "unreviewed",
            "notice": "Local unauthenticated operator names; dispositions are not financial approval or resolution.",
            "active_version": state["active_version"],
            "versions": versions,
            "rules": state["rules"],
            "findings": state["findings"],
            "changes": state["changes"],
            "events": state["events"],
            "checks_current": state["checks_current"],
            "mappings_confirmed": state["mappings_confirmed"],
            "mapping_confirmation": state["mapping_confirmation"],
            "limits": LIMITS,
        }
        packet["sha256"] = hashlib.sha256(_canonical(packet)).hexdigest()
        return packet
