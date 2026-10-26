from __future__ import annotations

import hashlib
import io
import os
import stat
import zipfile
from copy import deepcopy
from datetime import date
from xml.etree import ElementTree

import pytest
from openpyxl import Workbook

from workiva_mcp import workbench as workbench_module
from workiva_mcp.workbench import Workbench


def xlsx(sheets: dict[str, dict[str, object]]) -> bytes:
    book = Workbook()
    book.remove(book.active)
    for title, cells in sheets.items():
        sheet = book.create_sheet(title)
        for coordinate, value in cells.items():
            sheet[coordinate] = value
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def with_cached_formula(content: bytes, cached: str, formula: str | None = None) -> bytes:
    """Set formula text and its cached value independently (openpyxl cannot do this)."""
    source = io.BytesIO(content)
    output = io.BytesIO()
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(output, "w") as outgoing:
        for info in incoming.infolist():
            data = incoming.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(data)
                namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                cell = next(item for item in root.iter(namespace + "c") if item.attrib["r"] == "A1")
                formula_node = cell.find(namespace + "f")
                value_node = cell.find(namespace + "v")
                assert formula_node is not None and value_node is not None
                if formula is not None:
                    formula_node.text = formula.removeprefix("=")
                value_node.text = cached
                data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            outgoing.writestr(info, data)
    return output.getvalue()


def rule(
    rule_id: str = "r",
    *,
    reported: str = "A1",
    supporting: list[dict] | None = None,
    tolerance: str = "0.00",
    decimal_places: int = 2,
) -> dict:
    return {
        "id": rule_id,
        "description": f"Independent check {rule_id}",
        "reported": {"sheet": "Data", "cell": reported},
        "supporting": supporting
        or [{"sheet": "Data", "cell": "A2", "sign": "add", "scale": "units"}],
        "unit": "USD",
        "decimal_places": decimal_places,
        "tolerance": tolerance,
    }


def configured(tmp_path, content: bytes, rules: list[dict]) -> tuple[Workbench, dict]:
    wb = Workbench(tmp_path / "workspace")
    state = wb.import_workbook(content, "source.xlsx", "Source", 0)
    state = wb.save_rules(rules, state["revision"])
    state = wb.confirm_mappings("Mapper", state["revision"])
    return wb, wb.run_checks(state["revision"])


def by_id(state: dict, rule_id: str) -> dict:
    return next(item for item in state["findings"] if item["rule_id"] == rule_id)


def review(wb: Workbench, state: dict, rule_id: str, note: str = "Checked locally") -> dict:
    finding = by_id(state, rule_id)
    return wb.review_finding(
        rule_id, "noted", note, "Reviewer", finding["evidence_digest"], state["revision"]
    )


def test_initial_state_is_draft_and_cannot_run(tmp_path):
    wb = Workbench(tmp_path / "workspace")
    state = wb.state()
    assert state["revision"] == 0
    assert state["findings"] == []
    assert state["summary"] == {"passed": 0, "failed": 0, "not_evaluated": 0, "reopened": 0}
    assert state["mappings_confirmed"] is False
    assert state["checks_current"] is False
    assert state["demo_version"] is None
    with pytest.raises(ValueError, match="confirmed"):
        wb.run_checks(0)


def test_demo_v1_requires_confirmation_then_has_two_passes_and_two_unevaluated(tmp_path):
    wb = Workbench(tmp_path / "workspace")
    draft = wb.load_demo(1, 0)
    assert draft["demo_version"] == 1
    assert draft["mappings_confirmed"] is False
    assert draft["checks_current"] is False
    with pytest.raises(ValueError, match="confirmed"):
        wb.run_checks(draft["revision"])
    confirmed = wb.confirm_mappings("Mapper", draft["revision"])
    state = wb.run_checks(confirmed["revision"])
    assert state["summary"] == {"passed": 2, "failed": 0, "not_evaluated": 2, "reopened": 0}
    assert {item["rule_id"]: item["status"] for item in state["findings"]} == {
        "cash": "passed",
        "unaffected": "passed",
        "formula": "not_evaluated",
        "missing": "not_evaluated",
    }
    cash = by_id(state, "cash")
    assert (cash["reported"], cash["expected"], cash["difference"]) == (
        "1000000.00",
        "1000000.00",
        "0.00",
    )


def test_demo_v2_reopens_changed_evidence_but_retains_unaffected_review(tmp_path):
    wb = Workbench(tmp_path / "workspace")
    state = wb.load_demo(1, 0)
    state = wb.confirm_mappings("Mapper", state["revision"])
    state = wb.run_checks(state["revision"])
    for rule_id in ("cash", "formula", "unaffected"):
        state = review(wb, state, rule_id)
    state = wb.load_demo(2, state["revision"])
    assert state["demo_version"] == 2
    cash = by_id(state, "cash")
    assert (cash["status"], cash["difference"], cash["review_status"]) == (
        "failed",
        "5000.00",
        "reopened",
    )
    assert by_id(state, "formula")["review_status"] == "reopened"
    assert by_id(state, "unaffected")["review_status"] == "reviewed"
    assert state["summary"]["reopened"] == 2
    assert state["mappings_confirmed"] is False
    assert state["checks_current"] is False
    with pytest.raises(ValueError, match="current"):
        review(wb, state, "cash")
    state = wb.confirm_mappings("Mapper", state["revision"])
    rerun = wb.run_checks(state["revision"])
    assert rerun["summary"]["reopened"] == 2
    assert by_id(rerun, "cash")["review_status"] == "reopened"


def test_stale_revision_and_wrong_digest_are_atomic(tmp_path):
    wb, state = configured(tmp_path, xlsx({"Data": {"A1": 1, "A2": 1}}), [rule()])
    before = deepcopy(wb.state())
    with pytest.raises(ValueError, match="Stale"):
        wb.review_finding("r", "noted", "note", "Reviewer", by_id(state, "r")["evidence_digest"], 0)
    assert wb.state() == before
    with pytest.raises(ValueError, match="evidence"):
        wb.review_finding("r", "noted", "note", "Reviewer", "0" * 64, state["revision"])
    assert wb.state() == before


def test_restart_preserves_history_and_exported_evidence_snapshots(tmp_path):
    root = tmp_path / "workspace"
    wb, state = configured(tmp_path, xlsx({"Data": {"A1": 2, "A2": 2}}), [rule()])
    state = review(wb, state, "r")
    event_count = len(state["events"])
    reopened = Workbench(root)
    assert reopened.state() == state
    packet = reopened.export_packet()
    assert len(packet["events"]) == event_count
    assert packet["versions"][0]["evidence"]["Data"]["A1"]["value"] == 2
    assert packet["rules"] == state["rules"]
    assert packet["findings"] == state["findings"]
    assert packet["status"] == "unreviewed"
    assert "approval" in packet["notice"].lower() and "not" in packet["notice"].lower()


def test_rule_edit_makes_results_stale_and_rerun_reopens_affected_review(tmp_path):
    content = xlsx({"Data": {"A1": 10, "A2": 10, "A3": 10}})
    wb, state = configured(tmp_path, content, [rule()])
    state = review(wb, state, "r")
    changed = rule(supporting=[{"sheet": "Data", "cell": "A3", "sign": "add"}])
    draft = wb.save_rules([changed], state["revision"])
    assert draft["checks_current"] is False
    with pytest.raises(ValueError, match="current|confirmed|stale"):
        review(wb, draft, "r", "must not apply to old result")
    confirmed = wb.confirm_mappings("Mapper", draft["revision"])
    rerun = wb.run_checks(confirmed["revision"])
    assert by_id(rerun, "r")["review_status"] == "reopened"


def test_review_is_disposition_only_and_does_not_change_arithmetic(tmp_path):
    wb, state = configured(tmp_path, xlsx({"Data": {"A1": 12, "A2": 10}}), [rule()])
    before = by_id(state, "r")
    state = review(wb, state, "r")
    after = by_id(state, "r")
    assert (after["status"], after["reported"], after["expected"], after["difference"]) == (
        before["status"],
        before["reported"],
        before["expected"],
        before["difference"],
    )
    assert after["review"]["disposition"] == "noted"
    assert not any(key in after["review"] for key in ("approved", "resolved", "signature"))


def test_formula_uses_cached_value_and_raw_formula_is_part_of_evidence(tmp_path):
    original = xlsx({"Data": {"A1": "=A2", "A2": 7}})
    first = with_cached_formula(original, "7")
    wb, state = configured(tmp_path, first, [rule()])
    assert by_id(state, "r")["status"] == "passed"
    state = review(wb, state, "r")
    second = with_cached_formula(original, "7", "A2+0")
    state = wb.import_workbook(second, "second.xlsx", "Second", state["revision"])
    finding = by_id(state, "r")
    assert finding["status"] == "passed"
    assert finding["review_status"] == "reopened"


def test_preview_pages_retained_snapshot_without_mutating_workspace(tmp_path):
    content = with_cached_formula(
        xlsx({"First": {"A1": "=B1", "B1": 7, "K30": "distant"}, "Second": {"C3": 9}}),
        "7",
    )
    wb = Workbench(tmp_path / "workspace")
    state = wb.import_workbook(content, "source.xlsx", "Source", 0)
    before_state = deepcopy(wb.state())
    before_database = hashlib.sha256(wb.path.read_bytes()).hexdigest()

    page = wb.preview(state["revision"])
    assert page["revision"] == state["revision"]
    assert page["active_version"] == state["active_version"]
    assert page["sheets"] == ["First", "Second"]
    assert page["sheet"] == "First"
    assert (page["total_rows"], page["total_columns"]) == (30, 11)
    assert len(page["rows"]) == 20 and all(len(row) == 8 for row in page["rows"])
    assert page["rows"][0][0] == {
        "cell": "A1",
        "raw_value": "=B1",
        "value": 7,
        "error": None,
    }
    assert page["rows"][1][0] is None

    distant = wb.preview(state["revision"], "First", 25, 9)
    assert distant["rows"][5][2]["cell"] == "K30"
    second = wb.preview(state["revision"], "Second")
    assert second["rows"][2][2]["value"] == 9
    assert wb.state() == before_state
    assert hashlib.sha256(wb.path.read_bytes()).hexdigest() == before_database


def test_preview_validates_revision_sheet_and_excel_bounds(tmp_path):
    wb = Workbench(tmp_path / "workspace")
    with pytest.raises(ValueError, match="active workbook"):
        wb.preview(0)
    state = wb.import_workbook(xlsx({"Data": {"A1": 1}}), "source.xlsx", "Source", 0)
    with pytest.raises(ValueError, match="Stale"):
        wb.preview(0)
    with pytest.raises(ValueError, match="Unknown sheet"):
        wb.preview(state["revision"], "Missing")
    for kwargs in (
        {"expected_revision": True},
        {"expected_revision": state["revision"], "start_row": 0},
        {"expected_revision": state["revision"], "start_row": 1_048_577},
        {"expected_revision": state["revision"], "start_column": 0},
        {"expected_revision": state["revision"], "start_column": 16_385},
        {"expected_revision": state["revision"], "sheet": 1},
    ):
        with pytest.raises(ValueError, match="Invalid"):
            wb.preview(**kwargs)

    edge = wb.preview(state["revision"], start_row=1_048_576, start_column=16_384)
    assert (edge["start_row"], edge["start_column"]) == (1_048_557, 16_377)
    assert edge["rows"][-1][-1] is None


@pytest.mark.parametrize("bad", [None, "", True, date(2026, 1, 1), "#DIV/0!"])
def test_blank_boolean_date_and_error_are_not_coerced_to_zero(tmp_path, bad):
    wb, state = configured(tmp_path, xlsx({"Data": {"A1": bad, "A2": 0}}), [rule()])
    finding = by_id(state, "r")
    assert finding["status"] == "not_evaluated"
    assert finding["reported"] is None
    assert finding["reason"]


@pytest.mark.parametrize(
    ("reported", "support", "spec", "expected", "difference", "status"),
    [
        (
            1,
            [600, 400],
            {"reported_scale": "millions", "support_scale": "thousands"},
            "1000000.00",
            "0.00",
            "passed",
        ),
        (-2.345, [-2.344], {"places": 2, "tolerance": "0.00"}, "-2.34", "-0.01", "failed"),
        (101, [100], {"tolerance": "1.00"}, "100.00", "1.00", "passed"),
        (98.99, [100], {"tolerance": "1.00"}, "100.00", "-1.01", "failed"),
    ],
)
def test_numeric_calculation_has_signed_aggregate_scaling_rounding_and_symmetric_tolerance(
    tmp_path, reported, support, spec, expected, difference, status
):
    cells = {"A1": reported, **{f"A{i + 2}": value for i, value in enumerate(support)}}
    mappings = [
        {
            "sheet": "Data",
            "cell": f"A{i + 2}",
            "sign": "add",
            "scale": spec.get("support_scale", "units"),
        }
        for i in range(len(support))
    ]
    current_rule = rule(
        supporting=mappings,
        tolerance=spec.get("tolerance", "0.00"),
        decimal_places=spec.get("places", 2),
    )
    current_rule["reported"]["scale"] = spec.get("reported_scale", "units")
    wb, state = configured(tmp_path, xlsx({"Data": cells}), [current_rule])
    finding = by_id(state, "r")
    assert finding["expected"] == expected
    assert finding["difference"] == difference
    assert finding["status"] == status


def test_subtract_sign_is_applied_in_aggregate(tmp_path):
    mappings = [
        {"sheet": "Data", "cell": "A2", "sign": "add"},
        {"sheet": "Data", "cell": "A3", "sign": "subtract"},
    ]
    wb, state = configured(
        tmp_path, xlsx({"Data": {"A1": 60, "A2": 100, "A3": 40}}), [rule(supporting=mappings)]
    )
    finding = by_id(state, "r")
    assert (finding["expected"], finding["difference"], finding["status"]) == (
        "60.00",
        "0.00",
        "passed",
    )


def add_zip_entry(content: bytes, name: str, payload: bytes) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as incoming, zipfile.ZipFile(output, "w") as outgoing:
        for info in incoming.infolist():
            outgoing.writestr(info, incoming.read(info.filename))
        outgoing.writestr(name, payload)
    return output.getvalue()


@pytest.mark.parametrize("content", [b"", b"not a zip", b"PK\x03\x04truncated"])
def test_rejects_invalid_archives_without_state_change(tmp_path, content):
    wb = Workbench(tmp_path / "workspace")
    with pytest.raises(ValueError):
        wb.import_workbook(content, "bad.xlsx", "Bad", 0)
    assert wb.state()["revision"] == 0


@pytest.mark.parametrize("entry", ["xl/externalLinks/externalLink1.xml", "xl/vbaProject.bin"])
def test_rejects_external_links_and_vba(tmp_path, entry):
    wb = Workbench(tmp_path / "workspace")
    hostile = add_zip_entry(xlsx({"Data": {"A1": 1}}), entry, b"payload")
    with pytest.raises(ValueError, match="links|macros"):
        wb.import_workbook(hostile, "hostile.xlsx", "Hostile", 0)


def test_rejects_duplicate_archive_names(tmp_path):
    base = xlsx({"Data": {"A1": 1}})
    with pytest.warns(UserWarning, match="Duplicate name"):
        hostile = add_zip_entry(base, "xl/workbook.xml", b"duplicate")
    wb = Workbench(tmp_path / "workspace")
    with pytest.raises(ValueError, match="duplicate|archive"):
        wb.import_workbook(hostile, "duplicate.xlsx", "Duplicate", 0)


def test_rejects_absurd_dimension_and_sparse_far_away_cell(tmp_path):
    book = xlsx({"Data": {"A1": 1, "XFD1048576": 2}})
    wb = Workbench(tmp_path / "workspace")
    with pytest.raises(ValueError, match="dimension|cell|limit"):
        wb.import_workbook(book, "sparse.xlsx", "Sparse", 0)


def test_compressed_and_expanded_limits_are_enforced_cheaply(tmp_path, monkeypatch):
    content = xlsx({"Data": {"A1": "x" * 100}})
    wb = Workbench(tmp_path / "workspace")
    monkeypatch.setitem(workbench_module.LIMITS, "compressed_bytes", len(content) - 1)
    with pytest.raises(ValueError, match="compressed"):
        wb.import_workbook(content, "large.xlsx", "Large", 0)
    monkeypatch.setitem(workbench_module.LIMITS, "compressed_bytes", len(content) + 1)
    monkeypatch.setitem(workbench_module.LIMITS, "expanded_bytes", 1)
    with pytest.raises(ValueError, match="expanded"):
        wb.import_workbook(content, "expanded.xlsx", "Expanded", 0)


def test_import_retains_original_hash_does_not_mutate_input_and_sets_private_permissions(tmp_path):
    root = tmp_path / "workspace"
    content = xlsx({"Data": {"A1": 1}})
    original = bytes(content)
    wb = Workbench(root)
    state = wb.import_workbook(content, "source.xlsx", "Source", 0)
    assert content == original
    assert state["workbooks"][0]["sha256"] == hashlib.sha256(original).hexdigest()
    assert stat.S_IMODE(os.stat(root).st_mode) == 0o700
    assert stat.S_IMODE(os.stat(root / "workbench.sqlite3").st_mode) == 0o600


def test_unknown_sheet_cannot_be_confirmed(tmp_path):
    wb = Workbench(tmp_path / "workspace")
    state = wb.import_workbook(xlsx({"Data": {"A1": 1}}), "source.xlsx", "Source", 0)
    missing = rule()
    missing["reported"]["sheet"] = "Missing"
    state = wb.save_rules([missing], state["revision"])
    with pytest.raises(ValueError, match="unknown sheet"):
        wb.confirm_mappings("Mapper", state["revision"])


@pytest.mark.parametrize(
    "mutator",
    [
        lambda rules: rules.append(deepcopy(rules[0])),
        lambda rules: rules[0]["reported"].update(scale=[]),
        lambda rules: rules[0]["supporting"][0].update(scale=["units"]),
    ],
)
def test_malformed_scales_and_duplicate_rule_ids_are_rejected(tmp_path, mutator):
    wb = Workbench(tmp_path / "workspace")
    rules = [rule()]
    mutator(rules)
    with pytest.raises(ValueError):
        wb.save_rules(rules, 0)


def test_blank_reviewer_and_note_are_rejected(tmp_path):
    wb, state = configured(tmp_path, xlsx({"Data": {"A1": 1, "A2": 1}}), [rule()])
    digest = by_id(state, "r")["evidence_digest"]
    with pytest.raises(ValueError):
        wb.review_finding("r", "noted", "note", "", digest, state["revision"])
    with pytest.raises(ValueError):
        wb.review_finding("r", "noted", "", "Reviewer", digest, state["revision"])


def test_demo_cannot_replace_real_workspace_and_v1_is_atomic_on_stale_revision(tmp_path):
    wb = Workbench(tmp_path / "real")
    state = wb.import_workbook(xlsx({"Data": {"A1": 1}}), "real.xlsx", "Real", 0)
    before = deepcopy(state)
    with pytest.raises(ValueError, match="empty"):
        wb.load_demo(1, state["revision"])
    assert wb.state() == before

    empty = Workbench(tmp_path / "empty")
    with pytest.raises(ValueError, match="Stale"):
        empty.load_demo(1, 1)
    assert empty.state()["revision"] == 0
    assert empty.state()["workbooks"] == []


def test_packet_with_any_unreviewed_finding_has_no_approval_or_signature_claim(tmp_path):
    wb, state = configured(tmp_path, xlsx({"Data": {"A1": 1, "A2": 1}}), [rule()])
    packet = wb.export_packet()
    assert packet["status"] == "unreviewed"
    assert "signature" not in packet
    assert "approval" not in packet
    assert state["findings"][0]["review_status"] == "unreviewed"


@pytest.mark.parametrize(
    "part,payload",
    [
        ("custom/metadata.xml", b"<root>" + b"<x/>" * 500 + b"</root>"),
        ("custom/data.xml", b"<a>" * 40 + b"</a>" * 40),
        (
            "custom/links.rels",
            b"<Relationships>" + b'<Relationship Id="duplicate"/>' * 2 + b"</Relationships>",
        ),
    ],
)
def test_package_wide_xml_limits_precede_openpyxl(tmp_path, monkeypatch, part, payload):
    content = add_zip_entry(xlsx({"Data": {"A1": 1}}), part, payload)
    monkeypatch.setitem(workbench_module.LIMITS, "xml_elements", 400)
    monkeypatch.setattr(
        workbench_module,
        "load_workbook",
        lambda *a, **k: pytest.fail("Unbounded input reached parser"),
    )
    with pytest.raises(ValueError, match="complexity|relationships"):
        Workbench(tmp_path / "limited").import_workbook(content, "input.xlsx", "Input", 0)


def test_editing_one_rule_preserves_unaffected_disposition(tmp_path):
    wb, state = configured(
        tmp_path,
        xlsx({"Data": {"A1": 10, "A2": 10, "A3": 10}}),
        [rule("changed"), rule("untouched")],
    )
    state = review(wb, state, "changed")
    state = review(wb, state, "untouched")
    changed = deepcopy(state["rules"])
    changed[0]["supporting"][0]["cell"] = "A3"
    state = wb.save_rules(changed, state["revision"])
    state = wb.confirm_mappings("Mapper", state["revision"])
    state = wb.run_checks(state["revision"])
    assert by_id(state, "changed")["review_status"] == "reopened"
    assert by_id(state, "untouched")["review_status"] == "reviewed"
    assert state["summary"]["reopened"] == 1
